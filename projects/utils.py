import os
import zipfile
import re
import csv
from pathlib import Path
from pathlib import PurePosixPath, PureWindowsPath
from django.core.exceptions import ValidationError

EXECUTABLE_EXTENSIONS = {'.exe', '.bat', '.cmd', '.sh', '.msi', '.dll', '.so', '.dylib', '.com', '.scr', '.vbs'}
MAX_ZIP_SIZE = 50 * 1024 * 1024  # 50 MB
MAX_EXTRACTED_SIZE = 250 * 1024 * 1024
MAX_ZIP_ENTRIES = 10000
SRS_EXTENSIONS = {'.txt', '.pdf', '.docx', '.csv'}

def _validate_zip_members(zip_ref):
    entries = zip_ref.infolist()
    if len(entries) > MAX_ZIP_ENTRIES or sum(info.file_size for info in entries) > MAX_EXTRACTED_SIZE:
        raise ValidationError('ZIP contents exceed the allowed extraction limits.')
    for file_info in entries:
        name = file_info.filename
        posix_path = PurePosixPath(name.replace('\\', '/'))
        windows_path = PureWindowsPath(name)
        if '\x00' in name or posix_path.is_absolute() or windows_path.is_absolute() or windows_path.drive or '..' in posix_path.parts:
            raise ValidationError(f"Unsafe file path detected in ZIP: {name}")
        if posix_path.suffix.lower() in EXECUTABLE_EXTENSIONS:
            raise ValidationError(f"ZIP contains forbidden executable file type: {name}")

def validate_zip_upload(upload):
    if Path(upload.name).suffix.lower() != '.zip':
        raise ValidationError('Project source must be uploaded as a ZIP file.')
    if upload.size > MAX_ZIP_SIZE:
        raise ValidationError('Uploaded ZIP file exceeds the maximum allowed size of 50 MB.')
    upload.seek(0)
    try:
        with zipfile.ZipFile(upload) as zip_ref:
            _validate_zip_members(zip_ref)
    except (zipfile.BadZipFile, OSError, ValueError) as exc:
        raise ValidationError('The uploaded file is not a valid ZIP archive.') from exc
    finally:
        upload.seek(0)

def validate_srs_upload(upload):
    if Path(upload.name).suffix.lower() not in SRS_EXTENSIONS:
        raise ValidationError('SRS must be a TXT, PDF, DOCX, or CSV file.')
    if upload.size > MAX_ZIP_SIZE:
        raise ValidationError('SRS document exceeds the maximum allowed size of 50 MB.')

def validate_and_extract_zip(zip_file_path, target_dir):
    """
    Validates ZIP file for safety:
    - Check size <= 50MB
    - Check directory traversal (../)
    - Check executable files
    Extracts safely to target_dir.
    """
    if os.path.getsize(zip_file_path) > MAX_ZIP_SIZE:
        raise ValidationError("Uploaded ZIP file exceeds the maximum allowed size of 50 MB.")

    os.makedirs(target_dir, exist_ok=True)

    with zipfile.ZipFile(zip_file_path, 'r') as zip_ref:
        _validate_zip_members(zip_ref)

        # Safe extraction
        zip_ref.extractall(target_dir)

    return target_dir

def parse_csv_requirements(file_path):
    """
    Extracts requirements list from CSV file.
    Supports named headers (req_id, description, category, priority)
    or positional rows (e.g. REQ-001, System shall..., Authentication, High).
    """
    requirements = []
    content = ""
    for enc in ['utf-8-sig', 'utf-8', 'latin-1']:
        try:
            with open(file_path, 'r', encoding=enc) as f:
                content = f.read()
            break
        except Exception:
            continue

    if not content:
        return requirements

    lines = [l for l in content.splitlines() if l.strip()]
    if not lines:
        return requirements

    reader = list(csv.reader(lines))
    if not reader:
        return requirements

    header = [c.strip().lower() for c in reader[0]]
    id_keys = {'req_id', 'id', 'requirement_id', 'req id', 'requirement id', 'code', 'req'}
    desc_keys = {'description', 'desc', 'requirement', 'req_desc', 'statement', 'details', 'title'}
    cat_keys = {'category', 'cat', 'module', 'type', 'group'}
    prio_keys = {'priority', 'prio', 'severity', 'level'}

    id_idx = next((i for i, h in enumerate(header) if h in id_keys), None)
    desc_idx = next((i for i, h in enumerate(header) if h in desc_keys), None)
    cat_idx = next((i for i, h in enumerate(header) if h in cat_keys), None)
    prio_idx = next((i for i, h in enumerate(header) if h in prio_keys), None)

    has_named_header = (id_idx is not None or desc_idx is not None)
    data_rows = reader[1:] if has_named_header else reader

    req_pattern = re.compile(r'^(REQ[-_]?\d+|FR[-_]?\d+|\d+\.\d+)', re.IGNORECASE)
    count = 1

    for row in data_rows:
        if not row or not any(str(c).strip() for c in row):
            continue

        req_id = ''
        desc = ''
        category = 'General'
        priority = 'High'

        if has_named_header:
            if id_idx is not None and id_idx < len(row):
                req_id = row[id_idx].strip()
            if desc_idx is not None and desc_idx < len(row):
                desc = row[desc_idx].strip()
            if cat_idx is not None and cat_idx < len(row) and row[cat_idx].strip():
                category = row[cat_idx].strip()
            if prio_idx is not None and prio_idx < len(row) and row[prio_idx].strip():
                raw_prio = row[prio_idx].strip().capitalize()
                if raw_prio in ['High', 'Medium', 'Low']:
                    priority = raw_prio
        else:
            first_col = row[0].strip()
            if req_pattern.match(first_col):
                req_id = first_col
                desc = row[1].strip() if len(row) > 1 else ''
                if len(row) > 2 and row[2].strip():
                    category = row[2].strip()
                if len(row) > 3 and row[3].strip():
                    raw_prio = row[3].strip().capitalize()
                    if raw_prio in ['High', 'Medium', 'Low']:
                        priority = raw_prio
            else:
                desc = first_col
                if len(row) > 1 and row[1].strip():
                    category = row[1].strip()

        if not desc and req_id:
            desc = req_id
            req_id = f"REQ-{count:03d}"
        elif not req_id and desc:
            id_match = req_pattern.match(desc)
            if id_match:
                req_id = id_match.group(1).upper()
            else:
                req_id = f"REQ-{count:03d}"

        if desc:
            if not category or category == 'General':
                category = infer_category(desc)
            if not priority:
                priority = 'High' if 'must' in desc.lower() or 'shall' in desc.lower() else 'Medium'

            requirements.append({
                'req_id': req_id.upper(),
                'description': desc,
                'category': category,
                'priority': priority
            })
            count += 1

    return requirements

def parse_srs_file(file_path):
    """
    Extracts requirements list from TXT, DOCX, PDF, or CSV file.
    Returns list of dicts: [{'req_id': 'REQ-001', 'description': '...', 'category': '...', 'priority': 'High'}]
    """
    text = ""
    ext = os.path.splitext(file_path)[-1].lower()

    if ext == '.csv':
        csv_reqs = parse_csv_requirements(file_path)
        if csv_reqs:
            return csv_reqs
        # If CSV reader produced empty results, fall through to text reader fallback
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            text = f.read()
    elif ext == '.txt':
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            text = f.read()
    elif ext == '.docx':
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs]
            tables = ['\t'.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
            text = "\n".join(paragraphs + tables)
        except Exception as e:
            text = ""
    elif ext == '.pdf':
        try:
            import pypdf
            reader = pypdf.PdfReader(file_path)
            text = "\n".join([page.extract_text() or "" for page in reader.pages])
        except Exception as e:
            text = ""

    requirements = []
    req_pattern = re.compile(r'\b(FR[-_]?\d+)\b', re.IGNORECASE)
    section_starts = list(re.finditer(r'\b7\.\s*FUNCTIONAL\s+REQUIREMENTS', text, re.IGNORECASE))
    section_start = section_starts[-1].start() if section_starts else text.lower().find('functional requirements')
    section_end_match = re.search(r'\b8\.\s*NON[\s-]*FUNCTIONAL\s+REQUIREMENTS', text[section_start:], re.IGNORECASE) if section_start >= 0 else None
    section_end = section_start + section_end_match.start() if section_end_match else -1
    requirement_text = text[section_start:section_end] if section_start >= 0 and section_end > section_start else text
    explicit_ids = list(req_pattern.finditer(requirement_text))
    seen_ids = set()
    for index, match in enumerate(explicit_ids):
        req_id = match.group(1).upper()
        if req_id in seen_ids:
            continue
        end = explicit_ids[index + 1].start() if index + 1 < len(explicit_ids) else len(text)
        desc = re.sub(r'\s+', ' ', requirement_text[match.end():end]).strip(' :-.')
        priority_match = re.search(r'\s+(High|Medium|Low)\s*$', desc, re.IGNORECASE)
        priority = priority_match.group(1).capitalize() if priority_match else ('High' if re.search(r'\b(must|shall)\b', desc, re.IGNORECASE) else 'Medium')
        if priority_match:
            desc = desc[:priority_match.start()].strip()
        if not re.search(r'\b(shall|must|should)\b', desc, re.IGNORECASE):
            continue
        requirements.append({'req_id': req_id, 'description': desc, 'category': infer_category(desc), 'priority': priority})
        seen_ids.add(req_id)

    lines = [re.sub(r'\s+', ' ', line).strip() for line in text.splitlines() if line.strip()]
    next_id = 1
    for line in lines:
        explicit_req = re.match(r'^(REQ[-_]?\d+)\s*[:\-\.]?\s*(.+)$', line, re.IGNORECASE)
        if explicit_req:
            req_id, desc = explicit_req.group(1).upper(), explicit_req.group(2).strip()
            if req_id not in seen_ids and re.search(r'\b(shall|must|should)\b', desc, re.IGNORECASE):
                requirements.append({'req_id': req_id, 'description': desc, 'category': infer_category(desc), 'priority': 'High' if re.search(r'\b(must|shall)\b', desc, re.IGNORECASE) else 'Medium'})
                seen_ids.add(req_id)
            continue
        if ext == '.pdf' or req_pattern.search(line) or not re.search(r'\b(system|user)\s+shall\b|\b(must|should)\s+(allow|support|provide|include|store|accept|display|generate)\b', line, re.IGNORECASE):
            continue
        while f'REQ-{next_id:03d}' in seen_ids:
            next_id += 1
        req_id = f'REQ-{next_id:03d}'
        requirements.append({'req_id': req_id, 'description': line, 'category': infer_category(line), 'priority': 'High' if re.search(r'\b(must|shall)\b', line, re.IGNORECASE) else 'Medium'})
        seen_ids.add(req_id)
        next_id += 1

    return requirements

def infer_category(text):
    text_lower = text.lower()
    if any(k in text_lower for k in ['login', 'auth', 'user', 'role', 'password', 'permission']):
        return 'Authentication'
    elif any(k in text_lower for k in ['db', 'database', 'sqlite', 'model', 'data', 'store']):
        return 'Database'
    elif any(k in text_lower for k in ['api', 'rest', 'endpoint', 'post', 'get', 'json']):
        return 'API'
    elif any(k in text_lower for k in ['upload', 'file', 'zip', 'extract']):
        return 'File Handling'
    elif any(k in text_lower for k in ['test', 'unittest', 'pytest']):
        return 'Testing'
    elif any(k in text_lower for k in ['report', 'export', 'score', 'summary']):
        return 'Reporting'
    return 'General'
